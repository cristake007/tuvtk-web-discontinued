from __future__ import annotations

import csv
import io
from typing import Any
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from .domain import parse_excel_dates_from_row

TABULAR_EXTENSIONS = {".csv", ".xlsx", ".xls"}


class UpdateWordPressFileError(ValueError):
    pass


def _display_value(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _detect_csv_delimiter(sample: str) -> str:
    try:
        return csv.Sniffer().sniff(sample, delimiters=",;\t|").delimiter
    except csv.Error:
        return ";" if sample.count(";") > sample.count(",") else ","


def _read_csv_rows(data: bytes) -> list[list[Any]]:
    decoded: str | None = None
    for encoding in ("utf-8-sig", "utf-8", "cp1250", "latin-1"):
        try:
            decoded = data.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    if decoded is None:
        raise UpdateWordPressFileError("Fișierul CSV nu poate fi decodat.")

    delimiter = _detect_csv_delimiter(decoded[:8192])
    return [list(row) for row in csv.reader(io.StringIO(decoded), delimiter=delimiter)]


def _read_xlsx_rows(data: bytes) -> list[list[Any]]:
    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        sheet = workbook.active
        return [list(row) for row in sheet.iter_rows(values_only=True)]
    finally:
        workbook.close()


def _read_xls_rows(data: bytes) -> list[list[Any]]:
    try:
        import xlrd
    except ModuleNotFoundError as exc:
        raise UpdateWordPressFileError(
            "Suportul pentru fișiere .xls nu este instalat în backend. "
            "Reconstruiește imaginea backend sau folosește CSV/XLSX."
        ) from exc

    try:
        workbook = xlrd.open_workbook(file_contents=data)
    except xlrd.XLRDError as exc:
        raise UpdateWordPressFileError("Fișierul .xls încărcat nu poate fi citit.") from exc

    sheet = workbook.sheet_by_index(0)
    return [sheet.row_values(index) for index in range(sheet.nrows)]


def parse_course_date_rows(data: bytes, extension: str) -> list[dict[str, Any]]:
    readers = {
        ".csv": _read_csv_rows,
        ".xlsx": _read_xlsx_rows,
        ".xls": _read_xls_rows,
    }
    reader = readers.get(extension.casefold())
    if reader is None:
        raise UpdateWordPressFileError("Formatul fișierului nu este acceptat.")

    try:
        rows = reader(data)
    except UpdateWordPressFileError:
        raise
    except (
        BadZipFile,
        InvalidFileException,
        KeyError,
        OSError,
        UnicodeError,
        ValueError,
        csv.Error,
    ) as exc:
        raise UpdateWordPressFileError("Fișierul încărcat nu poate fi citit.") from exc

    if not rows:
        raise UpdateWordPressFileError("Fișierul nu conține date.")

    header_row = rows[0]
    normalized_headers = {
        str(value or "").strip().casefold(): index for index, value in enumerate(header_row)
    }
    title_index = normalized_headers.get("title")
    permalink_index = normalized_headers.get("permalink")
    if title_index is None or permalink_index is None:
        raise UpdateWordPressFileError(
            "Fișierul trebuie să conțină coloanele Title și Permalink."
        )

    parsed_rows: list[dict[str, Any]] = []
    for row_index, row in enumerate(rows[1:]):

        def cell(index: int | None) -> Any:
            if index is None or index >= len(row):
                return None
            return row[index]

        title = _display_value(cell(title_index))
        permalink = _display_value(cell(permalink_index))
        row_values = {
            str(header_row[index]): (row[index] if index < len(row) else None)
            for index in range(len(header_row))
        }
        excel_dates = parse_excel_dates_from_row(row_values)

        parsed_rows.append(
            {
                "row_index": row_index,
                "title": title,
                "permalink": permalink,
                "excel_dates": excel_dates,
            }
        )

    return parsed_rows
