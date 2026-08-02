import calendar
import csv
import io
import re
import unicodedata
from collections.abc import Iterable
from dataclasses import dataclass
from datetime import datetime
from typing import Any
from xml.dom import minidom
from xml.etree import ElementTree as ET
from zipfile import BadZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.utils.exceptions import InvalidFileException

MAX_COURSES = 10_000
MAX_HEADER_SCAN_ROWS = 25
EXCEL_FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r")

HEADER_ALIASES = {
    "title": ("title", "titlu", "nume curs"),
    "duration": ("durata curs", "duration", "durata"),
    "permalink": ("permalink", "link", "url"),
    "investment": ("investitie", "investment", "pret", "price"),
}
REQUIRED_HEADERS = ("title", "duration", "permalink")

XML_HEADER_ALIASES = {
    "title": ("title", "titlu", "nume curs"),
    "permalink": ("permalink", "link", "url"),
}
XML_REQUIRED_HEADERS = ("title", "permalink")
MONTH_COLUMN_ALIASES = {
    calendar.month_name[index].casefold(): index for index in range(1, 13)
} | {f"luna {index}": index for index in range(1, 13)}


class CourseFileError(ValueError):
    pass


@dataclass(frozen=True)
class CourseInput:
    title: str
    permalink: str
    duration_label: str
    investment: str
    duration: int
    original_order: int


def _normalize_header(value: Any) -> str:
    normalized = unicodedata.normalize("NFKD", str(value or ""))
    ascii_value = normalized.encode("ascii", "ignore").decode("ascii")
    return " ".join(ascii_value.strip().casefold().split())


def _display_value(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _neutralize_excel_formula(value: Any) -> Any:
    if not isinstance(value, str) or not value:
        return value

    first_non_space = value.lstrip(" ")[:1]
    if value.startswith(("\t", "\r")) or first_non_space in EXCEL_FORMULA_PREFIXES:
        return f"'{value}"
    return value


def _header_map(row: list[Any]) -> dict[str, int]:
    return {
        normalized: index
        for index, value in enumerate(row)
        if (normalized := _normalize_header(value))
    }


def _header_score(
    headers: dict[str, int],
    aliases: dict[str, tuple[str, ...]],
    required: tuple[str, ...],
) -> int:
    return sum(any(alias in headers for alias in aliases[name]) for name in required)


def _select_header_row(
    rows: list[list[Any]],
    aliases: dict[str, tuple[str, ...]] = HEADER_ALIASES,
    required: tuple[str, ...] = REQUIRED_HEADERS,
) -> tuple[int, dict[str, int]]:
    best_index: int | None = None
    best_headers: dict[str, int] = {}
    best_score = -1

    for index, row in enumerate(rows[:MAX_HEADER_SCAN_ROWS]):
        headers = _header_map(row)
        if not headers:
            continue

        score = _header_score(headers, aliases, required)
        if score > best_score:
            best_index = index
            best_headers = headers
            best_score = score

        if score == len(required):
            return index, headers

    if best_index is None:
        raise CourseFileError("Fișierul nu conține un rând de antet valid.")

    return best_index, best_headers


def _has_complete_header(
    rows: list[list[Any]],
    aliases: dict[str, tuple[str, ...]] = HEADER_ALIASES,
    required: tuple[str, ...] = REQUIRED_HEADERS,
) -> bool:
    try:
        _, headers = _select_header_row(rows, aliases, required)
    except CourseFileError:
        return False
    return _header_score(headers, aliases, required) == len(required)


def _parse_csv_rows(decoded: str, delimiter: str) -> list[list[Any]]:
    rows = [list(row) for row in csv.reader(io.StringIO(decoded), delimiter=delimiter)]
    parsed_rows: list[list[Any]] = []

    for row in rows:
        # Some legacy exports quote the complete record instead of individual fields,
        # for example: "Title@""Durata Curs""@investitie@Permalink".
        # The first CSV pass correctly removes the outer quotes but leaves one cell,
        # so parse that cell once more with the same delimiter.
        if len(row) == 1 and delimiter in str(row[0]):
            nested_row = next(csv.reader([str(row[0])], delimiter=delimiter))
            if len(nested_row) > 1:
                row = list(nested_row)
        parsed_rows.append(row)

    return parsed_rows


def _read_csv(
    data: bytes,
    aliases: dict[str, tuple[str, ...]] = HEADER_ALIASES,
    required: tuple[str, ...] = REQUIRED_HEADERS,
) -> list[list[Any]]:
    decoded: str | None = None
    for encoding in ("utf-8-sig", "utf-8", "cp1250", "latin-1"):
        try:
            decoded = data.decode(encoding)
            break
        except UnicodeDecodeError:
            continue

    if decoded is None:
        raise CourseFileError("Fișierul CSV nu poate fi decodat.")

    lines = decoded.splitlines()
    first_content_index = next(
        (index for index, line in enumerate(lines) if line.strip()),
        None,
    )
    if first_content_index is not None:
        separator_directive = re.fullmatch(
            r"\s*sep=(.)\s*",
            lines[first_content_index],
            flags=re.IGNORECASE,
        )
        if separator_directive is not None:
            delimiter = separator_directive.group(1)
            del lines[first_content_index]
            return _parse_csv_rows("\n".join(lines), delimiter)

    delimiters: list[str] = []
    sample = decoded[:8192]
    try:
        detected = csv.Sniffer().sniff(sample, delimiters=",;\t|@").delimiter
        delimiters.append(detected)
    except csv.Error:
        pass

    for delimiter in (";", ",", "\t", "|", "@"):
        if delimiter not in delimiters:
            delimiters.append(delimiter)

    fallback_rows: list[list[Any]] | None = None
    for delimiter in delimiters:
        rows = _parse_csv_rows(decoded, delimiter)
        if fallback_rows is None:
            fallback_rows = rows
        if _has_complete_header(rows, aliases, required):
            return rows

    return fallback_rows or []


def _read_xlsx(data: bytes) -> list[list[Any]]:
    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        sheet = workbook.active
        return [list(row) for row in sheet.iter_rows(values_only=True)]
    finally:
        workbook.close()


def _read_xls(data: bytes) -> list[list[Any]]:
    try:
        import xlrd
    except ModuleNotFoundError as exc:
        raise CourseFileError(
            "Suportul pentru fișiere .xls nu este instalat în backend. "
            "Reconstruiește imaginea backend sau folosește CSV/XLSX."
        ) from exc

    try:
        workbook = xlrd.open_workbook(file_contents=data)
    except xlrd.XLRDError as exc:
        raise CourseFileError("Fișierul .xls încărcat nu poate fi citit.") from exc

    sheet = workbook.sheet_by_index(0)
    return [sheet.row_values(index) for index in range(sheet.nrows)]


def _read_tabular_rows(
    data: bytes,
    extension: str,
    aliases: dict[str, tuple[str, ...]] = HEADER_ALIASES,
    required: tuple[str, ...] = REQUIRED_HEADERS,
) -> list[list[Any]]:
    normalized_extension = extension.casefold()
    if normalized_extension not in {".csv", ".xlsx", ".xls"}:
        raise CourseFileError("Formatul fișierului nu este acceptat.")

    try:
        if normalized_extension == ".csv":
            rows = _read_csv(data, aliases, required)
        elif normalized_extension == ".xlsx":
            rows = _read_xlsx(data)
        else:
            rows = _read_xls(data)
    except CourseFileError:
        raise
    except (
        BadZipFile,
        InvalidFileException,
        KeyError,
        OSError,
        UnicodeError,
        ValueError,
        csv.Error,
    ) as exc:
        raise CourseFileError("Fișierul încărcat nu poate fi citit.") from exc

    if not rows:
        raise CourseFileError("Fișierul nu conține date.")

    return rows


def parse_course_file(data: bytes, extension: str) -> list[CourseInput]:
    rows = _read_tabular_rows(data, extension)
    header_row_index, headers = _select_header_row(rows)

    def find_column(name: str, *, required: bool = True) -> int | None:
        for alias in HEADER_ALIASES[name]:
            if alias in headers:
                return headers[alias]
        if required:
            expected = HEADER_ALIASES[name][0]
            raise CourseFileError(f'Coloana obligatorie „{expected}” lipsește.')
        return None

    title_index = find_column("title")
    duration_index = find_column("duration")
    permalink_index = find_column("permalink")
    investment_index = find_column("investment", required=False)
    courses: list[CourseInput] = []

    for original_order, row in enumerate(rows[header_row_index + 1 :]):
        def cell(index: int | None, current_row=row) -> Any:
            if index is None or index >= len(current_row):
                return None
            return current_row[index]

        title = _display_value(cell(title_index))
        duration_label = _display_value(cell(duration_index))
        if not title or not duration_label:
            continue

        duration_match = re.search(r"\d+", duration_label)
        if duration_match is None:
            raise CourseFileError(
                f'Valoarea „{duration_label}” din coloana „Durata Curs” nu este validă.'
            )

        duration = int(duration_match.group())
        if not 1 <= duration <= 366:
            raise CourseFileError("Durata cursului trebuie să fie între 1 și 366 de zile.")

        courses.append(
            CourseInput(
                title=title,
                permalink=_display_value(cell(permalink_index)),
                duration_label=duration_label,
                investment=_display_value(cell(investment_index)),
                duration=duration,
                original_order=original_order,
            )
        )
        if len(courses) > MAX_COURSES:
            raise CourseFileError(f"Fișierul poate conține cel mult {MAX_COURSES} de cursuri.")

    if not courses:
        raise CourseFileError("Fișierul nu conține cursuri valide.")

    return courses


def create_excel_export(
    schedule: Iterable[dict[str, Any]],
    holidays: list[str],
) -> bytes:
    workbook = Workbook()
    schedule_sheet = workbook.active
    schedule_sheet.title = "Schedule"
    month_names = [
        "January",
        "February",
        "March",
        "April",
        "May",
        "June",
        "July",
        "August",
        "September",
        "October",
        "November",
        "December",
    ]
    headers = ["Title", "Permalink", "Durata Curs", "investitie", *month_names]
    schedule_sheet.append(headers)

    rows_by_order: dict[int, dict[str, Any]] = {}
    order: list[int] = []
    sorted_schedule = sorted(schedule, key=lambda value: value.get("original_order", 0))
    for fallback_order, item in enumerate(sorted_schedule):
        course_order = int(item.get("original_order", fallback_order))
        if course_order not in rows_by_order:
            rows_by_order[course_order] = {
                "Title": str(item["title"]),
                "Permalink": item.get("permalink", ""),
                "Durata Curs": item.get("duration_label", ""),
                "investitie": item.get("investment", ""),
                **{month: "" for month in month_names},
            }
            order.append(course_order)
        rows_by_order[course_order][month_names[int(item["month"]) - 1]] = item[
            "date_range"
        ]

    for course_order in order:
        row = rows_by_order[course_order]
        schedule_sheet.append(
            [_neutralize_excel_formula(row[header]) for header in headers]
        )

    if holidays:
        holidays_sheet = workbook.create_sheet("Holidays")
        holidays_sheet.append(["Holiday Date"])
        for holiday in holidays:
            holidays_sheet.append([_neutralize_excel_formula(holiday)])

    for sheet in workbook.worksheets:
        for column_index, column in enumerate(sheet.iter_cols(), start=1):
            width = max(len(str(cell.value or "")) for cell in column) + 2
            sheet.column_dimensions[get_column_letter(column_index)].width = min(width, 80)

    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def parse_xml_events(data: bytes, extension: str) -> list[dict[str, str]]:
    rows = _read_tabular_rows(data, extension, XML_HEADER_ALIASES, XML_REQUIRED_HEADERS)
    header_row_index, headers = _select_header_row(rows, XML_HEADER_ALIASES, XML_REQUIRED_HEADERS)

    def find_column(name: str) -> int:
        for alias in XML_HEADER_ALIASES[name]:
            if alias in headers:
                return headers[alias]
        expected = XML_HEADER_ALIASES[name][0]
        raise CourseFileError(f'Coloana obligatorie „{expected}” lipsește.')

    title_index = find_column("title")
    permalink_index = find_column("permalink")
    month_columns = sorted(
        index for header, index in headers.items() if header in MONTH_COLUMN_ALIASES
    )
    if not month_columns:
        raise CourseFileError(
            "Fișierul nu conține coloane de dată recunoscute "
            "(January-December sau Luna 1-Luna 12)."
        )

    events: list[dict[str, str]] = []
    for row in rows[header_row_index + 1 :]:
        def cell(index: int, current_row=row) -> Any:
            return current_row[index] if index < len(current_row) else None

        title = _display_value(cell(title_index))
        if not title:
            continue
        permalink = _display_value(cell(permalink_index))
        for month_index in month_columns:
            date_value = _display_value(cell(month_index))
            if date_value:
                events.append(
                    {
                        "course_name": title,
                        "date_range": date_value,
                        "permalink": permalink,
                    }
                )
        if len(events) > MAX_COURSES * 3:
            raise CourseFileError(f"Fișierul poate conține cel mult {MAX_COURSES} de cursuri.")

    if not events:
        raise CourseFileError("Fișierul nu conține date valide despre cursuri.")

    return events


def create_xml_export(events: list[dict[str, str]], start_post_id: int) -> str:
    def parse_date_range(date_range: str) -> tuple[str, str]:
        normalized = (date_range or "").strip()

        single_day = re.match(r"^(\d{1,2})\.(\d{1,2})\.(\d{4})$", normalized)
        if single_day:
            day, month, parsed_year = single_day.groups()
            iso_date = f"{parsed_year}-{month.zfill(2)}-{day.zfill(2)}"
            return iso_date, iso_date

        multi_day_same_month = re.match(
            r"^(\d{1,2})\s*-\s*(\d{1,2})\.(\d{1,2})\.(\d{4})$",
            normalized,
        )
        if multi_day_same_month:
            start_day, end_day, month, parsed_year = multi_day_same_month.groups()
            return (
                f"{parsed_year}-{month.zfill(2)}-{start_day.zfill(2)}",
                f"{parsed_year}-{month.zfill(2)}-{end_day.zfill(2)}",
            )

        multi_day_with_month = re.match(
            r"^(\d{1,2})\.(\d{1,2})\s*-\s*(\d{1,2})\.(\d{1,2})\.(\d{4})$",
            normalized,
        )
        if multi_day_with_month:
            (
                start_day,
                start_month,
                end_day,
                end_month,
                parsed_year,
            ) = multi_day_with_month.groups()
            return (
                f"{parsed_year}-{start_month.zfill(2)}-{start_day.zfill(2)}",
                f"{parsed_year}-{end_month.zfill(2)}-{end_day.zfill(2)}",
            )

        raise CourseFileError(f"Format de dată neacceptat: {date_range}")

    root = ET.Element("events")
    courses: dict[str, list[dict[str, str]]] = {}
    for event in events:
        courses.setdefault(event["course_name"], []).append(event)

    event_id = int(start_post_id)

    for course_name, course_events in courses.items():
        for period_index, event in enumerate(course_events, start=1):
            event_id += 1
            start_date, end_date = parse_date_range(str(event["date_range"]).strip())
            start_day_seconds = 8 * 3600
            end_day_seconds = 18 * 3600
            start_datetime = f"{start_date} 08:00 AM"
            end_datetime = f"{end_date} 06:00 PM"

            item = ET.SubElement(root, "item")
            ET.SubElement(item, "ID").text = str(event_id)
            ET.SubElement(item, "title").text = course_name
            ET.SubElement(item, "content").text = ""

            post = ET.SubElement(item, "post")
            ET.SubElement(post, "ID").text = str(event_id)
            ET.SubElement(post, "post_author").text = "5"
            ET.SubElement(post, "post_date").text = f"{start_date} 00:00:00"
            ET.SubElement(post, "post_date_gmt").text = f"{start_date} 00:00:00"
            ET.SubElement(post, "post_title").text = course_name
            ET.SubElement(post, "post_status").text = "draft"

            meta = ET.SubElement(item, "meta")
            ET.SubElement(meta, "mec_more_info_title").text = f"perioada {period_index}"
            ET.SubElement(meta, "mec_read_more").text = event.get("permalink", "")
            ET.SubElement(meta, "mec_color").text = ""
            ET.SubElement(meta, "mec_location_id").text = "1"
            ET.SubElement(meta, "mec_organizer_id").text = "1"
            ET.SubElement(meta, "mec_allday").text = "1"
            ET.SubElement(meta, "mec_start_date").text = start_date
            ET.SubElement(meta, "mec_start_time_hour").text = "8"
            ET.SubElement(meta, "mec_start_time_minutes").text = "00"
            ET.SubElement(meta, "mec_start_time_ampm").text = "AM"
            ET.SubElement(meta, "mec_start_day_seconds").text = str(start_day_seconds)
            ET.SubElement(meta, "mec_start_datetime").text = start_datetime
            ET.SubElement(meta, "mec_end_date").text = end_date
            ET.SubElement(meta, "mec_end_time_hour").text = "6"
            ET.SubElement(meta, "mec_end_time_minutes").text = "00"
            ET.SubElement(meta, "mec_end_time_ampm").text = "PM"
            ET.SubElement(meta, "mec_end_day_seconds").text = str(end_day_seconds)
            ET.SubElement(meta, "mec_end_datetime").text = end_datetime
            ET.SubElement(meta, "mec_repeat_status").text = "0"

            mec_date = ET.SubElement(meta, "mec_date")
            date_start = ET.SubElement(mec_date, "start")
            ET.SubElement(date_start, "date").text = start_date
            ET.SubElement(date_start, "hour").text = "8"
            ET.SubElement(date_start, "minutes").text = "00"
            ET.SubElement(date_start, "ampm").text = "AM"
            date_end = ET.SubElement(mec_date, "end")
            ET.SubElement(date_end, "date").text = end_date
            ET.SubElement(date_end, "hour").text = "6"
            ET.SubElement(date_end, "minutes").text = "00"
            ET.SubElement(date_end, "ampm").text = "PM"
            ET.SubElement(mec_date, "allday").text = "1"

            mec_block = ET.SubElement(item, "mec")
            ET.SubElement(mec_block, "id").text = ""
            ET.SubElement(mec_block, "post_id").text = str(event_id)
            ET.SubElement(mec_block, "start").text = start_date
            ET.SubElement(mec_block, "end").text = end_date
            ET.SubElement(mec_block, "repeat").text = "0"
            ET.SubElement(mec_block, "time_start").text = str(start_day_seconds)
            ET.SubElement(mec_block, "time_end").text = str(end_day_seconds)

            time_block = ET.SubElement(item, "time")
            ET.SubElement(time_block, "start").text = "All Day"
            ET.SubElement(time_block, "end").text = ""
            ET.SubElement(time_block, "start_raw").text = "8:00 am"
            ET.SubElement(time_block, "end_raw").text = "6:00 pm"
            ET.SubElement(time_block, "start_timestamp").text = str(
                int(datetime.strptime(start_date, "%Y-%m-%d").timestamp()) + start_day_seconds
            )
            ET.SubElement(time_block, "end_timestamp").text = str(
                int(datetime.strptime(end_date, "%Y-%m-%d").timestamp()) + end_day_seconds
            )

    return minidom.parseString(ET.tostring(root)).toprettyxml(indent="    ")
