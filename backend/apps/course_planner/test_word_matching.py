import base64
import io
import tempfile

from django.contrib.auth import get_user_model
from django.contrib.auth.models import Group
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase, override_settings
from docx import Document

from apps.core.models import ApplicationDefinition

from .word_matching import (
    ScheduleRow,
    apply_word_matches,
    build_word_course_rows,
    confident_match,
    normalize_title,
    parse_schedule_file,
    score_schedule_matches,
)


def build_word_document(rows: list[tuple[str, str, str, str, str, str]]) -> bytes:
    document = Document()
    table = document.add_table(rows=len(rows), cols=6)
    for row_index, cells in enumerate(rows):
        for column_index, value in enumerate(cells):
            table.rows[row_index].cells[column_index].text = value
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


class NormalizeTitleTests(TestCase):
    def test_strips_diacritics_and_punctuation(self):
        self.assertEqual(
            normalize_title("Auditor ISO 9001:2015 — Șef echipă"),
            "auditor iso 9001 2015 sef echipa",
        )


class ScheduleMatchingTests(TestCase):
    def test_exact_normalized_title_short_circuits_scoring(self):
        schedule = [
            ScheduleRow(0, "Auditor ISO 9001", normalize_title("Auditor ISO 9001"), ("", "", "")),
            ScheduleRow(1, "Manager calitate", normalize_title("Manager calitate"), ("", "", "")),
        ]

        matches = score_schedule_matches("Auditor ISO 9001", schedule)

        self.assertEqual(len(matches), 1)
        self.assertTrue(matches[0].exact)
        self.assertEqual(matches[0].entry.row_index, 0)

    def test_confident_match_returns_none_below_threshold(self):
        schedule = [
            ScheduleRow(0, "Auditor ISO 9001", normalize_title("Auditor ISO 9001"), ("", "", "")),
        ]

        matches = score_schedule_matches("Complet neinrudit subiect", schedule)

        self.assertIsNone(confident_match("Complet neinrudit subiect", matches))

    def test_parse_schedule_file_reads_first_three_dates_in_calendar_order(self):
        schedule_rows = parse_schedule_file(
            b"Title;February;January\n"
            b"Auditor ISO 9001;15.02.2027;10.01.2027\n",
            ".csv",
        )

        self.assertEqual(len(schedule_rows), 1)
        row = schedule_rows[0]
        self.assertEqual(row.title, "Auditor ISO 9001")
        self.assertEqual(row.dates, ("10.01.2027", "15.02.2027", ""))


class WordDocumentMatchingTests(TestCase):
    def test_build_word_course_rows_skips_merged_header_row(self):
        document = Document()
        table = document.add_table(rows=2, cols=6)
        header_cells = table.rows[0].cells
        header_cells[0].merge(header_cells[-1])
        header_cells[0].text = "Ianuarie"
        for column_index, value in enumerate(
            ("Auditor ISO 9001", "2 zile", "350 EUR", "", "", "")
        ):
            table.rows[1].cells[column_index].text = value

        rows = build_word_course_rows(document)

        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0].title, "Auditor ISO 9001")

    def test_apply_word_matches_writes_dates_into_matched_rows(self):
        word_bytes = build_word_document(
            [
                ("Auditor ISO 9001", "2 zile", "350 EUR", "", "", ""),
                ("Curs necunoscut", "1 zi", "150 EUR", "", "", ""),
            ]
        )
        schedule = [
            ScheduleRow(0, "Auditor ISO 9001", normalize_title("Auditor ISO 9001"), ("11-12.01.2027", "", "")),
        ]

        output_bytes, matched_count, skipped_count = apply_word_matches(
            word_bytes, schedule, {0: 0}
        )

        self.assertEqual(matched_count, 1)
        self.assertEqual(skipped_count, 1)
        updated = Document(io.BytesIO(output_bytes))
        table = updated.tables[0]
        self.assertEqual(table.rows[0].cells[3].text, "11-12.01.2027")
        self.assertEqual(table.rows[1].cells[3].text, "")


class WordMatchApiTests(TestCase):
    def setUp(self):
        self.media_directory = tempfile.TemporaryDirectory()
        self.media_override = override_settings(MEDIA_ROOT=self.media_directory.name)
        self.media_override.enable()
        self.user = get_user_model().objects.create_user(
            username="word-matcher-user",
            password="test-password-123",
        )
        application, _ = ApplicationDefinition.objects.get_or_create(
            key="course-planner",
            defaults={"name": "Planificator cursuri"},
        )
        application.is_enabled = True
        application.is_access_controlled = True
        application.save(update_fields=("is_enabled", "is_access_controlled", "updated_at"))
        group = Group.objects.create(name="Word matcher test")
        application.allowed_groups.add(group)
        self.user.groups.add(group)
        self.client.force_login(self.user)

    def tearDown(self):
        self.media_override.disable()
        self.media_directory.cleanup()
        super().tearDown()

    def test_preview_matches_word_rows_against_schedule_file(self):
        word_file = SimpleUploadedFile(
            "cursuri.docx",
            build_word_document(
                [("Auditor ISO 9001", "2 zile", "350 EUR", "", "", "")]
            ),
            content_type=(
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
            ),
        )
        schedule_file = SimpleUploadedFile(
            "program.csv",
            b"Title;January\nAuditor ISO 9001;11-12.01.2027\n",
            content_type="text/csv",
        )

        response = self.client.post(
            "/api/v1/course-planner/word-match/preview",
            {"word_file": word_file, "schedule_file": schedule_file},
        )

        self.assertEqual(response.status_code, 200)
        payload = response.json()
        self.assertEqual(payload["matched_count"], 1)
        self.assertEqual(payload["skipped_count"], 0)
        self.assertEqual(payload["rows"][0]["status"], "matched")
        self.assertEqual(payload["rows"][0]["selected_row_index"], 0)
        self.assertIn("word_file_b64", payload)
        return payload

    def test_generate_applies_confirmed_matches_and_returns_docx(self):
        preview = self.test_preview_matches_word_rows_against_schedule_file()

        response = self.client.post(
            "/api/v1/course-planner/word-match/generate",
            data={
                "word_file_b64": preview["word_file_b64"],
                "schedule_options": preview["schedule_options"],
                "matches": [{"word_row_index": 0, "schedule_row_index": 0}],
            },
            content_type="application/json",
        )

        self.assertEqual(response.status_code, 200)
        self.assertTrue(
            response["Content-Type"].startswith(
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
            )
        )
        document = Document(io.BytesIO(response.content))
        self.assertEqual(document.tables[0].rows[0].cells[3].text, "11-12.01.2027")

    def test_generate_rejects_duplicate_schedule_row_indexes(self):
        response = self.client.post(
            "/api/v1/course-planner/word-match/generate",
            data={
                "word_file_b64": base64.b64encode(
                    build_word_document(
                        [("Auditor ISO 9001", "2 zile", "350 EUR", "", "", "")]
                    )
                ).decode("ascii"),
                "schedule_options": [
                    {"row_index": 0, "title": "Curs", "dates": ["", "", ""]},
                    {"row_index": 0, "title": "Curs 2", "dates": ["", "", ""]},
                ],
                "matches": [],
            },
            content_type="application/json",
        )

        self.assertEqual(response.status_code, 400)
        self.assertIn("indici duplicați", response.json()["detail"])

    def test_word_match_endpoints_require_application_access(self):
        user_without_permission = get_user_model().objects.create_user(
            username="no-word-matcher-access",
            password="test-password-123",
        )
        self.client.force_login(user_without_permission)

        preview_response = self.client.post(
            "/api/v1/course-planner/word-match/preview",
            {
                "word_file": SimpleUploadedFile("a.docx", b"x"),
                "schedule_file": SimpleUploadedFile("a.csv", b"x"),
            },
        )
        generate_response = self.client.post(
            "/api/v1/course-planner/word-match/generate",
            data={"word_file_b64": "", "schedule_options": [], "matches": []},
            content_type="application/json",
        )

        self.assertEqual(preview_response.status_code, 403)
        self.assertEqual(generate_response.status_code, 403)
