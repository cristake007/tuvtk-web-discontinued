import io
import tempfile
from datetime import date
from unittest.mock import patch

from django.contrib.auth import get_user_model
from django.contrib.auth.models import Permission
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase, override_settings
from openpyxl import load_workbook

from apps.public_holidays.models import PublicHoliday

from .files import create_excel_export
from .models import CoursePlanGeneration
from .scheduler import CourseScheduler


class CourseSchedulerTests(TestCase):
    def test_short_course_stays_inside_work_week(self):
        scheduler = CourseScheduler(2027, [])

        available = scheduler.get_available_start_days(1, 5)

        self.assertTrue(available)
        self.assertTrue(all(value.weekday() == 0 for value in available))

    def test_long_course_formats_cross_month_range(self):
        scheduler = CourseScheduler(2027, [])

        date_range = scheduler.format_date_range(date(2027, 2, 15), 18)

        self.assertEqual(date_range, "15.02-10.03.2027")


class CoursePlannerApiTests(TestCase):
    def setUp(self):
        self.media_directory = tempfile.TemporaryDirectory()
        self.media_override = override_settings(MEDIA_ROOT=self.media_directory.name)
        self.media_override.enable()
        self.user = get_user_model().objects.create_user(
            username="planner-user",
            password="test-password-123",
        )
        permission = Permission.objects.get(
            codename="use_course_planner",
            content_type__app_label="course_planner",
        )
        self.user.user_permissions.add(permission)
        self.client.force_login(self.user)

    def tearDown(self):
        self.media_override.disable()
        self.media_directory.cleanup()
        super().tearDown()

    def course_file(self, content: bytes | None = None) -> SimpleUploadedFile:
        return SimpleUploadedFile(
            "cursuri.csv",
            content
            or (
                b"Title;Durata Curs;investitie;Permalink\n"
                b"Auditor ISO 9001;2 zile;350 EUR;/curs/auditor-iso-9001\n"
            ),
            content_type="text/csv",
        )

    def generate_request(self, input_file: SimpleUploadedFile):
        return self.client.post(
            "/api/v1/course-planner/generate",
            {
                "input_file": input_file,
                "year": 2027,
                "months": "1",
                "randomness": 5,
            },
        )

    def test_generate_requires_authentication(self):
        self.client.logout()

        response = self.generate_request(self.course_file())

        self.assertEqual(response.status_code, 401)

    def test_authenticated_user_without_module_permission_receives_forbidden(self):
        user_without_permission = get_user_model().objects.create_user(
            username="no-planner-access",
            password="test-password-123",
        )
        self.client.force_login(user_without_permission)
        generation_id = "00000000-0000-0000-0000-000000000001"

        responses = [
            self.client.get("/api/v1/course-planner/settings"),
            self.client.get("/api/v1/course-planner/holidays?year=2028&months=3"),
            self.client.put(
                "/api/v1/course-planner/settings",
                data={"year": 2028, "months": [3], "randomness": 5},
                content_type="application/json",
            ),
            self.generate_request(self.course_file()),
            self.client.get("/api/v1/course-planner/history"),
            self.client.post(
                "/api/v1/course-planner/export",
                data={"generation_id": generation_id},
                content_type="application/json",
            ),
            self.client.get(
                f"/api/v1/course-planner/history/{generation_id}/files/excel"
            ),
        ]

        self.assertTrue(responses)
        for response in responses:
            self.assertEqual(response.status_code, 403)

    def test_holidays_lists_only_the_selected_planner_period(self):
        PublicHoliday.objects.create(
            date=date(2027, 1, 1),
            local_name="Anul Nou",
            source=PublicHoliday.Source.NAGER,
            is_internal=False,
        )
        PublicHoliday.objects.create(
            date=date(2027, 2, 2),
            local_name="Zi internă companie",
            source=PublicHoliday.Source.INTERNAL,
            is_internal=True,
        )
        PublicHoliday.objects.create(
            date=date(2027, 3, 8),
            local_name="Zi în afara selecției",
            source=PublicHoliday.Source.INTERNAL,
            is_internal=True,
        )

        response = self.client.get("/api/v1/course-planner/holidays?year=2027&months=1,2")

        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response.json(),
            [
                {
                    "date": "01.01.2027",
                    "name": "Anul Nou",
                    "source_label": "Zi legală",
                },
                {
                    "date": "02.02.2027",
                    "name": "Zi internă companie",
                    "source_label": "Zi internă",
                },
            ],
        )

    @patch(
        "apps.course_planner.views.get_holiday_dates",
        return_value=[date(2027, 1, 1)],
    )
    def test_generate_schedule_uses_backend_holidays_and_persists_history(
        self,
        get_holiday_dates,
    ):
        response = self.client.post(
            "/api/v1/course-planner/generate",
            {
                "input_file": self.course_file(),
                "year": 2027,
                "months": "1,2",
                "randomness": 5,
            },
        )

        self.assertEqual(response.status_code, 200)
        payload = response.json()
        self.assertEqual(payload["course_count"], 1)
        self.assertEqual(len(payload["schedule"]), 2)
        self.assertEqual(payload["schedule"][0]["title"], "Auditor ISO 9001")
        self.assertTrue(payload["generation"]["files"]["excel"]["available"])
        self.assertFalse(payload["generation"]["files"]["word"]["available"])
        self.assertFalse(payload["generation"]["files"]["xml"]["available"])
        get_holiday_dates.assert_called_once_with(year=2027, months=[1, 2])

        generation = CoursePlanGeneration.objects.get()
        self.assertEqual(generation.created_by, self.user)
        self.assertEqual(generation.source_filename, "cursuri.csv")
        self.assertEqual(generation.months, [1, 2])
        self.assertEqual(generation.holidays, ["01.01.2027"])
        self.assertEqual(len(generation.schedule), 2)
        self.assertTrue(generation.excel_file.storage.exists(generation.excel_file.name))

    def test_history_lists_generation_and_downloads_saved_excel(self):
        generate_response = self.generate_request(self.course_file())
        generation_id = generate_response.json()["generation"]["id"]

        history_response = self.client.get("/api/v1/course-planner/history")

        self.assertEqual(history_response.status_code, 200)
        history = history_response.json()
        self.assertEqual(history["count"], 1)
        self.assertEqual(history["results"][0]["id"], generation_id)
        self.assertEqual(history["results"][0]["created_by"], "planner-user")

        download_response = self.client.get(
            f"/api/v1/course-planner/history/{generation_id}/files/excel"
        )
        self.assertEqual(download_response.status_code, 200)
        self.assertTrue(
            download_response["Content-Type"].startswith(
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            )
        )
        workbook = load_workbook(
            io.BytesIO(b"".join(download_response.streaming_content)),
            read_only=True,
        )
        self.assertEqual(workbook["Schedule"]["A2"].value, "Auditor ISO 9001")
        workbook.close()

    def test_unavailable_history_artifact_returns_not_found(self):
        generate_response = self.generate_request(self.course_file())
        generation_id = generate_response.json()["generation"]["id"]

        response = self.client.get(
            f"/api/v1/course-planner/history/{generation_id}/files/word"
        )

        self.assertEqual(response.status_code, 404)

    def test_generate_csv_with_excel_separator_directive(self):
        response = self.generate_request(
            self.course_file(
                b"sep=;\n\n"
                b"Title;Durata Curs;investitie;Permalink\n"
                b"Auditor ISO 9001;2 zile;350 EUR;/curs/auditor-iso-9001\n"
            )
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()["schedule"][0]["title"], "Auditor ISO 9001")

    def test_generate_csv_with_header_after_preamble(self):
        response = self.generate_request(
            self.course_file(
                b"Export cursuri\n\n"
                b"Title;Durata Curs;investitie;Permalink\n"
                b"Auditor ISO 9001;2 zile;350 EUR;/curs/auditor-iso-9001\n"
            )
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()["schedule"][0]["title"], "Auditor ISO 9001")

    def test_generate_quoted_at_delimited_csv(self):
        response = self.generate_request(
            self.course_file(
                b'\xef\xbb\xbf"Title@""Durata Curs""@investitie@Permalink"\r\n'
                b'"Auditor ISO 9001@""2 zile""@""350 EUR""@'
                b'https://example.com/curs/auditor-iso-9001"\r\n'
            )
        )

        self.assertEqual(response.status_code, 200)
        payload = response.json()
        self.assertEqual(payload["course_count"], 1)
        self.assertEqual(payload["schedule"][0]["title"], "Auditor ISO 9001")
        self.assertEqual(payload["schedule"][0]["duration_label"], "2 zile")
        self.assertEqual(payload["schedule"][0]["investment"], "350 EUR")

    def test_invalid_input_columns_are_rejected(self):
        response = self.generate_request(self.course_file(b"Nume;Durata\nCurs;2\n"))

        self.assertEqual(response.status_code, 400)
        self.assertIn("Coloana obligatorie", response.json()["detail"])

    def test_settings_can_be_updated(self):
        response = self.client.put(
            "/api/v1/course-planner/settings",
            data={
                "year": 2028,
                "months": [3, 4],
                "randomness": 7,
            },
            content_type="application/json",
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()["months"], [3, 4])
        self.assertNotIn("holidays", response.json())

    @patch(
        "apps.course_planner.views.get_holiday_dates",
        return_value=[date(2027, 1, 1)],
    )
    def test_schedule_export_uses_saved_generation_data(self, get_holiday_dates):
        generate_response = self.generate_request(self.course_file())
        generation_id = generate_response.json()["generation"]["id"]

        response = self.client.post(
            "/api/v1/course-planner/export",
            data={"generation_id": generation_id},
            content_type="application/json",
        )

        self.assertEqual(response.status_code, 200)
        self.assertTrue(
            response["Content-Type"].startswith(
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            )
        )
        workbook = load_workbook(io.BytesIO(response.content), read_only=True)
        self.assertEqual(workbook["Schedule"]["A2"].value, "Auditor ISO 9001")
        self.assertEqual(workbook["Holidays"]["A2"].value, "01.01.2027")
        workbook.close()

    def test_excel_export_keeps_distinct_courses_with_duplicate_titles(self):
        excel_data = create_excel_export(
            [
                {
                    "title": "Auditor ISO 9001",
                    "permalink": "/curs/first",
                    "duration_label": "2 zile",
                    "investment": "350 EUR",
                    "date_range": "11-12.01.2027",
                    "month": 1,
                    "original_order": 0,
                },
                {
                    "title": "Auditor ISO 9001",
                    "permalink": "/curs/second",
                    "duration_label": "3 zile",
                    "investment": "450 EUR",
                    "date_range": "18-20.01.2027",
                    "month": 1,
                    "original_order": 1,
                },
            ],
            [],
        )

        workbook = load_workbook(io.BytesIO(excel_data), read_only=True)
        schedule = workbook["Schedule"]
        self.assertEqual(schedule.max_row, 3)
        self.assertEqual(schedule["A2"].value, "Auditor ISO 9001")
        self.assertEqual(schedule["A3"].value, "Auditor ISO 9001")
        self.assertEqual(schedule["B2"].value, "/curs/first")
        self.assertEqual(schedule["B3"].value, "/curs/second")
        workbook.close()
